"""Acme corpus: hand-written spec -> Haiku prose -> PDF/MD/DOCX/HTML files, checked fact by fact."""

import argparse
import html
import json
import re
import textwrap
import time
import unicodedata
from datetime import date as Date
from datetime import datetime
from pathlib import Path

import boto3
import docx

from engine import llm
from engine.parsers import parse

SPEC = Path(__file__).resolve().parent.parent / "data" / "acme" / "spec.json"
# Document id prefix -> (what the LLM writes, file suffix, the shape it should follow).
KINDS = {
    "INC": (
        "an internal incident report (postmortem)",
        ".pdf",
        "Sections: Summary, Impact, Timeline (one '- HH:MM UTC event' bullet per event, only "
        "times given here), Root cause, Resolution, Follow-up actions.",
    ),
    "CHG": (
        "an internal change record",
        ".md",
        "Sections: Description, Reason, Rollout, Rollback plan, Approval, Status.",
    ),
    "RB": (
        "an internal operations runbook",
        ".docx",
        "Sections: Purpose, When to use, Steps (one '- ' bullet per step), Escalation.",
    ),
    "STS": (
        "a public status page entry for merchants",
        ".html",
        "Short, plain language, no internal document ids or employee names. One '## ' section "
        "per update (for example Investigating, Identified, Monitoring, Resolved; Scheduled, "
        "Completed; or one section per service for a report), as fits the summary.",
    ),
}
DOC_ID = re.compile(r"^(INC|CHG|RB|STS)-\d{2,4}$")
QUESTION_TYPES = {"simple", "multi_hop", "multi_doc", "unanswerable", "contradictory"}
# Haiku allows 10 requests/min, shared with other jobs: one uncached call per 7 s leaves room.
PACE_SECONDS = 7
HAIKU_PRICE = (1.10, 5.50)  # USD per 1M input / output tokens, EU profile (ADR 0017)


# Spec


def load_spec(path: Path = SPEC) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def file_name(doc_id: str) -> str:
    """INC-301 -> acme-inc-301.pdf: unique per id, only [a-z0-9.-]."""
    return f"acme-{doc_id.lower()}{KINDS[doc_id.split('-')[0]][1]}"


def source(doc_id: str) -> str:
    """The document's source once `ingest` stores the uploaded file."""
    return f"upload:{file_name(doc_id)}"


def required_facts(doc: dict) -> list[str]:
    """Strings the rendered document must contain: its facts plus the ids it refers to."""
    return doc["facts"] + doc["refs"]


def validate(spec: dict) -> list[str]:
    """Everything that would make the labels lie; empty list = valid."""
    problems = []
    docs = {}
    for doc in spec["documents"]:
        doc_id = doc["id"]
        if not DOC_ID.match(doc_id):
            problems.append(f"{doc_id}: id must look like INC-301, CHG-1042, RB-03 or STS-101")
            continue
        if doc_id in docs:
            problems.append(f"{doc_id}: duplicate id")
        docs[doc_id] = doc
        try:
            Date.fromisoformat(doc["date"])
        except ValueError:
            problems.append(f"{doc_id}: date {doc['date']!r} is not YYYY-MM-DD")
        for name in doc["services"]:
            if name not in spec["company"]["services"]:
                problems.append(f"{doc_id}: unknown service {name}")
        if not doc["facts"]:
            problems.append(f"{doc_id}: no facts")
        for text in [doc["title"], doc["summary"], *doc["facts"]]:
            if not text.isascii() or "|" in text:
                problems.append(f"{doc_id}: non-ASCII or '|' in {text!r}")
    for doc in docs.values():
        for ref in doc["refs"]:
            if ref not in docs or ref == doc["id"]:
                problems.append(f"{doc['id']}: bad reference {ref}")

    contradictions = {c["id"]: c for c in spec["contradictions"]}
    for c in spec["contradictions"]:
        values = [_flat(claim["value"]) for claim in c["claims"]]
        if len(values) < 2 or len(set(values)) != len(values):
            problems.append(f"{c['id']}: needs at least 2 claims with different values")
        for claim in c["claims"]:
            doc = docs.get(claim["doc"])
            if doc is None:
                problems.append(f"{c['id']}: unknown document {claim['doc']}")
            elif not any(_flat(claim["value"]) in _flat(fact) for fact in doc["facts"]):
                problems.append(f"{c['id']}: {claim['value']!r} is not a fact of {claim['doc']}")

    gaps = {g["id"]: g for g in spec["gaps"]}
    for gap in spec["gaps"]:
        for doc in docs.values():
            found = _forbidden_in(" ".join([doc["summary"], *doc["facts"]]), gap["forbidden"])
            if found:
                problems.append(f"{gap['id']}: {doc['id']} spec mentions {found}")

    for number, q in enumerate(spec["questions"], start=1):
        where = f"question {number}"
        if q["type"] not in QUESTION_TYPES:
            problems.append(f"{where}: unknown type {q['type']}")
        elif q["type"] == "unanswerable":
            if q.get("gap") not in gaps or q.get("gold") or q.get("answer"):
                problems.append(f"{where}: unanswerable needs a known gap, no gold, no answer")
        elif q["type"] == "contradictory":
            if q.get("contradiction") not in contradictions:
                problems.append(f"{where}: unknown contradiction {q.get('contradiction')}")
        elif not q.get("answer") or not q.get("gold"):
            problems.append(f"{where}: needs an answer and gold documents")
        for doc_id in q.get("gold", []):
            if doc_id not in docs:
                problems.append(f"{where}: unknown gold document {doc_id}")
    return problems


# Checks


def _flat(text: str) -> str:
    """Case- and whitespace-insensitive form: PDF line wrapping and 'Approved by' still match."""
    return " ".join(text.split()).lower()


def missing_facts(text: str, facts: list[str]) -> list[str]:
    flat = _flat(text)
    return [fact for fact in facts if _flat(fact) not in flat]


def _forbidden_in(text: str, words: list[str]) -> list[str]:
    flat = _flat(text)
    return [word for word in words if _flat(word) in flat]


def problems_in(text: str, doc: dict, gaps: list[dict]) -> list[str]:
    """What is wrong with a generated document: required facts missing, gap words present."""
    words = [word for gap in gaps for word in gap["forbidden"]]
    return [f"missing {fact!r}" for fact in missing_facts(text, required_facts(doc))] + [
        f"must not mention {word!r}" for word in _forbidden_in(text, words)
    ]


# ponytail: ASCII, not latin-1. The PDF writer below uses the standard Helvetica encoding, and
# every format gets the same text, so a fact checked in one format holds in all of them.
_PUNCTUATION = str.maketrans(
    {
        "‘": "'", "’": "'", "‚": "'", "‛": "'",
        "“": '"', "”": '"', "„": '"',
        "‐": "-", "‑": "-", "‒": "-", "–": "-", "—": "-",
        "―": "-", "−": "-", "•": "-", "×": "x",
        "→": "->", "←": "<-", "≥": ">=", "≤": "<=",
    }
)  # fmt: skip


def sanitize(text: str) -> str:
    """Typographic punctuation -> ASCII; accents dropped (e -> e); anything else removed."""
    text = unicodedata.normalize("NFKD", text.translate(_PUNCTUATION))
    return text.encode("ascii", "ignore").decode("ascii")


# Generation


def prompts(doc: dict, spec: dict) -> tuple[str, str]:
    """System prompt (same for every document) and the user prompt for one document."""
    company = spec["company"]
    services = "\n".join(f"- {name}: {about}" for name, about in company["services"].items())
    people = "\n".join(f"- {name}: {role}" for name, role in company["people"].items())
    gaps = "\n".join(f"- {gap['what']}" for gap in spec["gaps"])
    system = (
        "You write documents for Acme Payments, a fictional company. The documents form a test "
        "corpus with known facts, so every fact must be exactly as given.\n\n"
        f"{company['about']}\n\nServices:\n{services}\n\nPeople:\n{people}\n\n"
        "Rules:\n"
        "- Write only the document body in Markdown: '## ' section headings, paragraphs and "
        "'- ' bullets. No title line, no tables, no bold or italics, no code fences.\n"
        "- Copy every required fact exactly as written (same words, digits and punctuation) at "
        "least once.\n"
        "- Do not add times, dates, numbers, people, document ids or causes beyond the ones "
        "given. Generic procedure text is fine.\n"
        "- Mention other documents only by their id; do not describe what they say.\n"
        "- Plain ASCII punctuation.\n"
        "- Acme does not know the following. Never mention, estimate or invent them:\n"
        f"{gaps}"
    )
    what, _, shape = KINDS[doc["id"].split("-")[0]]
    titles = {d["id"]: d["title"] for d in spec["documents"]}
    refs = "\n".join(f"- {ref} ({titles[ref]})" for ref in doc["refs"]) or "- none"
    facts = "\n".join(f"- {fact}" for fact in doc["facts"])
    user = (
        f"Write {what}, document {doc['id']}.\n"
        f"Title: {doc['title']}\nDate: {doc['date']}\nServices: {', '.join(doc['services'])}\n"
        f"What it is about: {doc['summary']}\n\n"
        f"Required facts (copy each exactly):\n{facts}\n\n"
        f"Documents to refer to by id:\n{refs}\n\n"
        f"Shape: {shape}\nLength: 150 to 400 words."
    )
    return system, user


class _Stats:
    def __init__(self) -> None:
        self.calls = self.cached = self.input_tokens = self.output_tokens = 0
        self.models: dict[str, int] = {}

    def cost(self) -> float:
        return (self.input_tokens * HAIKU_PRICE[0] + self.output_tokens * HAIKU_PRICE[1]) / 1e6


def _converse(system: str, messages: list[dict], stats: _Stats) -> dict:
    result = llm.converse(system, messages)
    stats.calls += 1
    if result["cached"]:
        stats.cached += 1
    else:
        stats.input_tokens += result["usage"]["inputTokens"]
        stats.output_tokens += result["usage"]["outputTokens"]
        stats.models[result["model"]] = stats.models.get(result["model"], 0) + 1
        time.sleep(PACE_SECONDS)
    return result["message"]


def generate(doc: dict, spec: dict, files: Path, stats: _Stats) -> list[str]:
    """Write one document; check the parsed file; one retry naming what was wrong.

    Returns the problems left after the retry (empty = the file in `files` is good). The check
    runs on the text `parse` extracts from the rendered file, so it covers both the LLM
    leaving a fact out and a renderer losing it.
    """
    system, user = prompts(doc, spec)
    messages = [{"role": "user", "content": [{"text": user}]}]
    for attempt in (1, 2):
        message = _converse(system, messages, stats)
        body = "".join(block.get("text", "") for block in message["content"])
        path = render(doc, body, files)
        problems = problems_in(parse(path, source(doc["id"])).text, doc, spec["gaps"])
        if not problems:
            return []
        print(f"{doc['id']} attempt {attempt}: {'; '.join(problems)}")
        feedback = (
            "The document has these problems: " + "; ".join(problems) + ". Rewrite the whole "
            "document body, copying every required fact exactly as written."
        )
        messages += [message, {"role": "user", "content": [{"text": feedback}]}]
    return problems


# Rendering


def blocks(body: str) -> list[tuple[str, str]]:
    """Markdown body -> [(kind, text)]: kind "h2".."h6", "li" or "p". Fences and ** dropped."""
    out = []
    for line in sanitize(body).splitlines():
        line = line.strip().replace("**", "")
        if not line or line.startswith("```"):
            continue
        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            # The title is ours (h1); a heading the model wrote is at least a section.
            out.append((f"h{min(max(level, 2), 6)}", line.lstrip("#").strip()))
        elif line.startswith(("- ", "* ")):
            out.append(("li", line[2:].strip()))
        else:
            out.append(("p", line))
    return out


def render(doc: dict, body: str, out_dir: Path) -> Path:
    """Write the document file. Title, id, date and the `Services:` line come from the spec,
    placed where `parsers.parse` finds them; the body is the LLM's text."""
    header = [
        ("h1", doc["title"]),
        ("p", f"Document: {doc['id']}"),
        ("p", f"Date: {doc['date']}"),
        ("p", f"Services: {', '.join(doc['services'])}"),
    ]
    content = header + blocks(body)
    path = out_dir / file_name(doc["id"])
    out_dir.mkdir(parents=True, exist_ok=True)
    suffix = path.suffix
    if suffix == ".md":
        path.write_text(_markdown(content), encoding="utf-8")
    elif suffix == ".html":
        path.write_text(_html(doc["title"], content), encoding="utf-8")
    elif suffix == ".docx":
        _docx(doc, content, path)
    else:
        path.write_bytes(_pdf(doc, content))
    return path


def _markdown(content: list[tuple[str, str]]) -> str:
    lines = []
    for kind, text in content:
        if kind == "li":
            lines.append("- " + text)
        elif kind == "p":
            lines.append(text)
        else:
            lines.append("#" * int(kind[1]) + " " + text)
    return "\n\n".join(lines) + "\n"


def _html(title: str, content: list[tuple[str, str]]) -> str:
    parts = []
    in_list = False
    for kind, text in content:
        if in_list and kind != "li":
            parts.append("</ul>")
            in_list = False
        if kind == "li" and not in_list:
            parts.append("<ul>")
            in_list = True
        tag = "li" if kind == "li" else kind
        parts.append(f"<{tag}>{html.escape(text)}</{tag}>")
    if in_list:
        parts.append("</ul>")
    body = "\n".join(parts)
    return (
        '<!doctype html>\n<html lang="en">\n<head>\n<meta charset="utf-8">\n'
        f"<title>{html.escape(title)}</title>\n</head>\n<body>\n{body}\n</body>\n</html>\n"
    )


def _docx(doc: dict, content: list[tuple[str, str]], path: Path) -> None:
    document = docx.Document()
    for kind, text in content:
        if kind == "h1":
            document.add_heading(text, level=0)  # the "Title" style: parse() reads it as "# "
        elif kind.startswith("h"):
            document.add_heading(text, level=int(kind[1]))
        elif kind == "li":
            document.add_paragraph(text, style="List Bullet")
        else:
            document.add_paragraph(text)
    props = document.core_properties
    props.title = doc["title"]
    props.author = "Acme Payments"
    props.created = datetime.fromisoformat(doc["date"])
    document.save(str(path))


# A4 in points; 10 pt Helvetica averages ~5 pt per character, so ~95 characters fill a line.
_PAGE_W, _PAGE_H, _MARGIN = 595, 842, 56
_STYLE = {"h1": ("F2", 16, 60), "li": ("F1", 10, 92), "p": ("F1", 10, 95)}  # font, size, chars


def _pdf(doc: dict, content: list[tuple[str, str]]) -> bytes:
    """A minimal multi-page PDF: one text line per `Tj`, Helvetica and Helvetica-Bold.

    No PDF library: lines are wrapped by character count (textwrap), and each page is a list
    of absolutely positioned lines. Title, author and creation date go in the Info dictionary,
    where pypdf reads them. ponytail: no headings in the extracted text and approximate
    wrapping; a layout engine (fpdf2) only if the files need to look typeset.
    """
    pages: list[list[str]] = [[]]
    y = _PAGE_H - _MARGIN
    for kind, text in content:
        font, size, width = _STYLE.get(kind, ("F2", 12, 80))  # h2..h6: bold 12 pt
        lines = textwrap.wrap(text, width)
        if kind == "li":
            lines = ["- " + lines[0]] + ["  " + line for line in lines[1:]]
        for line in lines:
            if y - size < _MARGIN:
                pages.append([])
                y = _PAGE_H - _MARGIN
            y -= size + 4
            pages[-1].append(f"BT /{font} {size} Tf {_MARGIN} {y} Td ({_pdf_text(line)}) Tj ET")
        y -= 6  # gap between blocks

    created = doc["date"].replace("-", "")
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [{}] /Count {} >>".format(
            " ".join(f"{6 + 2 * i} 0 R" for i in range(len(pages))), len(pages)
        ),
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold >>",
        f"<< /Title ({_pdf_text(doc['title'])}) /Author (Acme Payments) "
        f"/CreationDate (D:{created}000000+00'00') >>",
    ]
    for i, lines in enumerate(pages):
        stream = "\n".join(lines)
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 {_PAGE_W} {_PAGE_H}] "
            f"/Resources << /Font << /F1 3 0 R /F2 4 0 R >> >> /Contents {7 + 2 * i} 0 R >>"
        )
        objects.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")

    # The xref table lists the byte offset of every object, collected as they are appended.
    out = b"%PDF-1.4\n"
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n{body}\nendobj\n".encode("ascii")
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += "".join(f"{offset:010d} 00000 n \n" for offset in offsets).encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R /Info 5 0 R >>\n".encode()
    out += f"startxref\n{xref}\n%%EOF\n".encode()
    return out


def _pdf_text(text: str) -> str:
    """Escape a string for a PDF literal: backslash and parentheses."""
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


# Labels


def ground_truth(spec: dict) -> str:
    """data/acme/GROUND_TRUTH.md: documents, planted contradictions, gaps, question labels."""
    docs = spec["documents"]
    out = [
        "# Acme Payments corpus: ground truth",
        "",
        "Generated by `scripts/gen_acme.py` from `data/acme/spec.json` (`make acme-labels`); "
        "do not edit by hand. The files are in `data/acme/files/`; once loaded through "
        "`uploads/` each document's source is `upload:<file>`.",
        "",
        spec["company"]["about"],
        "",
        f"## Documents ({len(docs)})",
        "",
        "| File | Date | Services | Title | Summary |",
        "|---|---|---|---|---|",
    ]
    for doc in sorted(docs, key=lambda d: (d["date"], d["id"])):
        out.append(
            f"| `{file_name(doc['id'])}` | {doc['date']} | {', '.join(doc['services'])} | "
            f"{doc['title']} | {doc['summary']} |"
        )
    out += [
        "",
        f"## Planted contradictions ({len(spec['contradictions'])})",
        "",
        "Each value is a required fact of its document, checked verbatim in the generated file.",
        "",
        "| Id | What | Conflicting values |",
        "|---|---|---|",
    ]
    for c in spec["contradictions"]:
        values = "; ".join(f'{claim["doc"]}: "{claim["value"]}"' for claim in c["claims"])
        out.append(f"| {c['id']} | {c['what']} | {values} |")
    out += [
        "",
        f"## Gaps ({len(spec['gaps'])})",
        "",
        "Facts absent from every document: each generation prompt forbids them, and the words "
        "listed are checked to be absent from every generated file.",
        "",
        "| Id | Missing everywhere | Checked absent |",
        "|---|---|---|",
    ]
    for gap in spec["gaps"]:
        words = ", ".join(f"`{word}`" for word in gap["forbidden"]) or "-"
        out.append(f"| {gap['id']} | {gap['what']} | {words} |")
    rows = questions(spec)
    out += [
        "",
        f"## Questions ({len(rows)})",
        "",
        "Also in `data/acme/questions.jsonl` (the eval runner's input).",
        "",
        "| Id | Type | Question | Reference answer | Gold documents |",
        "|---|---|---|---|---|",
    ]
    for row in rows:
        gold = ", ".join(f"`{s.removeprefix('upload:')}`" for s in row["gold_sources"]) or "-"
        answer = row["answer"] if row["answer"] is not None else "(unanswerable)"
        out.append(f"| {row['id']} | {row['type']} | {row['question']} | {answer} | {gold} |")
    return "\n".join(out) + "\n"


def questions(spec: dict) -> list[dict]:
    """Eval questions in the runner's format. Contradictory answers and gold documents come
    from the contradiction they name, so the labels cannot drift from the documents."""
    contradictions = {c["id"]: c for c in spec["contradictions"]}
    rows = []
    for number, q in enumerate(spec["questions"], start=1):
        answer, gold = q.get("answer"), q.get("gold", [])
        if q["type"] == "contradictory":
            claims = contradictions[q["contradiction"]]["claims"]
            answer = "conflicting: " + " vs ".join(f"{c['value']} ({c['doc']})" for c in claims)
            gold = [c["doc"] for c in claims]
        elif q["type"] == "unanswerable":
            answer, gold = None, []
        rows.append(
            {
                "id": f"acme-{number:03d}",
                "question": q["question"],
                "type": q["type"],
                "answer": answer,
                "gold_sources": [source(doc_id) for doc_id in gold],
            }
        )
    return rows


def write_labels(spec: dict, out: Path) -> None:
    out.mkdir(parents=True, exist_ok=True)
    (out / "GROUND_TRUTH.md").write_text(ground_truth(spec), encoding="utf-8")
    lines = "".join(json.dumps(row, ensure_ascii=False) + "\n" for row in questions(spec))
    (out / "questions.jsonl").write_text(lines, encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--bucket", help="upload the corpus to s3://BUCKET/jobs/acme/")
    parser.add_argument("--out", type=Path, default=SPEC.parent, help="default: data/acme")
    parser.add_argument("--labels-only", action="store_true", help="labels only, no LLM")
    args = parser.parse_args()

    spec = load_spec()
    problems = validate(spec)
    if problems:
        raise SystemExit("invalid spec:\n" + "\n".join(problems))
    write_labels(spec, args.out)
    print(f"labels written to {args.out}")
    if args.labels_only:
        return
    if not args.bucket:
        raise SystemExit("--bucket is required to generate the corpus")

    files = args.out / "files"
    stats = _Stats()
    failed = {}
    for doc in spec["documents"]:
        problems = generate(doc, spec, files, stats)
        if problems:
            failed[doc["id"]] = problems
        print(f"{doc['id']}: {'FAILED' if problems else 'ok'}")
    print(
        f"{stats.calls} LLM calls ({stats.cached} cached), {stats.input_tokens} input + "
        f"{stats.output_tokens} output tokens, ${stats.cost():.3f}, models {stats.models}"
    )
    if failed:
        raise SystemExit("ground-truth check failed:\n" + json.dumps(failed, indent=2))

    # jobs/ is inert: uploads/ would start ingestion (that is `make acme-load`, after review).
    s3 = boto3.client("s3")
    for path in sorted(p for p in args.out.rglob("*") if p.is_file()):
        key = f"jobs/acme/{path.relative_to(args.out).as_posix()}"
        s3.upload_file(str(path), args.bucket, key)
    print(
        f"uploaded {len(spec['documents'])} documents and labels to s3://{args.bucket}/jobs/acme/"
    )


if __name__ == "__main__":
    main()
